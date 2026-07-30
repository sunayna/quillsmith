"""
Parses "23_24_Learning Standards.xlsx" into the same {grade: {subject:
[standard text]}} shape as build_reference.py produces from the 2025-26
workbook — but this workbook is laid out completely differently: one sheet
per subject, grades side by side as column groups (not one sheet per grade),
and no structural row labels ("Standard N", etc) to look for.

The one reliable signal across every sheet: row 2 always contains a "Grade N"
header in whatever column holds that grade's standard text — column
positions vary per sheet (and even per grade within a sheet, e.g. DL gives
each grade its own Skill+Grade pair), so those headers are located by
scanning row 2 rather than assumed fixed. Once a grade's column is known,
every non-empty cell below it (skipping the trailing thousands of blank
formatted rows past the real data) is a candidate standard text — the
"Skill" columns alongside are grouping labels only, not needed here since
this reference is used purely to test "is this text a real standard",
not to reproduce the workbook's own skill hierarchy.

Usage: python3 build_reference_2023_24.py
Writes: reference/standards_2023_24.json
"""
import json
import os
import re

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.join(os.path.dirname(__file__), "..", "..", "Term Data", "23_24_Learning Standards.xlsx")
XLSX_PATH = os.path.normpath(XLSX_PATH)
OUT_PATH = os.path.join(os.path.dirname(__file__), "standards_2023_24.json")

# Sheet name -> subject key, matching live-tree subject naming rather than
# the workbook's own abbreviated sheet names.
SHEET_TO_SUBJECT = {
    "DL": "Digital Literacy",
    "English": "English",
    "Expedition": "Expedition",
    "Hindi": "Hindi",
    "Math": "Mathematics",
    "Module": "Module",
    "French": "French (Third Language)",
    "Sanskrit": "Sanskrit (Third Language)",
}

GRADE_HEADER_RE = re.compile(r"Grade\s*(\d+)")
ARABIC_TO_ROMAN = {4: "IV", 5: "V", 6: "VI", 7: "VII"}


def find_grade_columns(header_row):
    """Returns {roman_grade: column_index} from row 2's cell values."""
    cols = {}
    for idx, val in enumerate(header_row):
        if not val:
            continue
        m = GRADE_HEADER_RE.search(str(val))
        if m:
            arabic = int(m.group(1))
            roman = ARABIC_TO_ROMAN.get(arabic)
            if roman:
                cols[roman] = idx
    return cols


def find_last_data_row(ws):
    """Sheets carry thousands of blank formatted rows past the real data —
    find the true last row with any content instead of trusting max_row."""
    last = 0
    for r in range(1, ws.max_row + 1):
        if any(c.value not in (None, "") for c in ws[r]):
            last = r
    return last


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    result = {}

    for sheet_name in wb.sheetnames:
        subject = SHEET_TO_SUBJECT.get(sheet_name)
        if not subject:
            print(f"⚠️  Unrecognized sheet {sheet_name!r} — skipping")
            continue

        ws = wb[sheet_name]
        header_row = [c.value for c in ws[2]]
        grade_cols = find_grade_columns(header_row)
        if not grade_cols:
            print(f"⚠️  No 'Grade N' header found in {sheet_name} row 2 — skipping")
            continue

        last_row = find_last_data_row(ws)

        for roman, col in grade_cols.items():
            texts = []
            for r in range(3, last_row + 1):
                val = ws.cell(row=r, column=col + 1).value
                if val is None:
                    continue
                text = str(val).strip()
                if text and text not in texts:
                    texts.append(text)
            result.setdefault(roman, {}).setdefault(subject, [])
            result[roman][subject].extend(t for t in texts if t not in result[roman][subject])

    json.dump(result, open(OUT_PATH, "w"), ensure_ascii=False, indent=2)

    print(f"✅ Wrote {OUT_PATH}")
    for grade, subjects in result.items():
        counts = {s: len(v) for s, v in subjects.items()}
        print(f"   {grade}: {counts}")


if __name__ == "__main__":
    main()
