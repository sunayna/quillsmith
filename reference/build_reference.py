"""
Builds a ground-truth reference of real Standard-level text per grade/subject
from the school's own "ASSESSMENT TREES" workbook — needed because reportbee's
own tree structure has no fully reliable signal for "this is the Standard"
versus "this is a Topic/Skill heading, an assessment instance (FA/SA), or a
rubric criterion beneath the real standard". A short_name convention
("S1", "S2", ...) looked like that signal at first but isn't: confirmed on an
Expedition standard where the real text ("Compares types of governments...")
sat one level below a node that also matched the S<n> pattern ("Governments:
Who Holds the Power?", itself just a topic heading).

Matching tree node names against this known-correct text is far more
reliable than any structural heuristic over reportbee's tree.

Usage: python3 build_reference.py "<path to ASSESSMENT TREES.xlsx>" [output.json]
"""
import json
import re
import sys

import openpyxl

STRUCTURAL = re.compile(r'^(skill|topic|standard|assessments|marks|weightage|lt\d+)', re.I)
STANDARD_ROW = re.compile(r'^Standard\s+\d+:?$')

# Sheet name -> grade label used elsewhere in this project (Roman numeral).
SHEET_TO_GRADE = {'G4': 'IV', 'G5': 'V', 'G6': 'VI', 'G7': 'VII'}


def parse_sheet(ws):
    """Returns {subject: [standard_text, ...]} for one grade sheet."""
    current_subject = None
    standards = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        b_empty = b is None or (isinstance(b, str) and not b.strip())

        if isinstance(a, str) and a.strip() and b_empty and not STRUCTURAL.match(a.strip()):
            current_subject = a.strip()
        elif isinstance(a, str) and STANDARD_ROW.match(a.strip()) and isinstance(b, str) and b.strip():
            if current_subject is None:
                continue
            standards.setdefault(current_subject, []).append(b.strip())

    return standards


def main():
    if len(sys.argv) < 2:
        print('usage: python3 build_reference.py "<ASSESSMENT TREES.xlsx>" [output.json]')
        sys.exit(1)

    xlsx_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else 'standards_reference.json'

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    reference = {}
    for sheet_name, grade in SHEET_TO_GRADE.items():
        if sheet_name not in wb.sheetnames:
            print(f'⚠️  Sheet {sheet_name} not found — skipping')
            continue
        standards = parse_sheet(wb[sheet_name])
        reference[grade] = standards
        total = sum(len(v) for v in standards.values())
        print(f'{sheet_name} -> Grade {grade}: {total} standards across {len(standards)} subjects')

    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(reference, f, ensure_ascii=False, indent=2)
    print(f'wrote {out_path}')


if __name__ == '__main__':
    main()
