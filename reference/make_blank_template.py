"""
Generates a blank "ASSESSMENT TREES"-style workbook in the exact shape
build_reference.py already knows how to parse — sheets named G4, G5, G6, G7
(Grade 4-7), where a row with text in column A and an EMPTY column B starts
a new Subject, and rows shaped "Standard N" / "Standard N:" in column A with
the real standard text in column B are the ones actually captured ("Skill"/
"Topic Name" rows in between are fine, they're just skipped — see
build_reference.py's STRUCTURAL/STANDARD_ROW patterns).

Meant to be handed to the school (or filled in directly) for a year that
doesn't have a reference workbook yet — filter_standards.py falls back to
reusing the closest existing year's reference until a dedicated one like
this gets filled in and parsed with build_reference.py.

Usage: python3 make_blank_template.py <year-label, e.g. 2026-27> [out-path]
"""
import os
import sys

import openpyxl

GRADES = ["G4", "G5", "G6", "G7"]
EXAMPLE_SUBJECTS_PER_GRADE = 3
EXAMPLE_STANDARDS_PER_SUBJECT = 3


def build_instructions_sheet(wb):
    ws = wb.active
    ws.title = "Instructions"
    lines = [
        "How to fill in this workbook",
        "",
        "One sheet per grade: G4, G5, G6, G7.",
        "",
        "Within a sheet:",
        "- A row with text in column A and column B left EMPTY starts a new",
        "  Subject (e.g. 'Mathematics', 'English', 'Expedition').",
        "- Rows labelled 'Skill', 'Topic Name', 'Skill/Topic' etc. in column A",
        "  are optional grouping labels — leave column B empty on those, they",
        "  get skipped.",
        "- A row with 'Standard 1', 'Standard 2', ... in column A and the",
        "  ACTUAL standard text in column B is what gets captured as that",
        "  subject's standard. Number sequentially within each subject.",
        "- Repeat: Subject row, then its Standard rows, then the next",
        "  Subject row.",
        "",
        "Each grade sheet already has a placeholder example shape (blank",
        "standard text for you to fill in) — replace the placeholder",
        "subject/standard names and add as many Subject/Standard blocks as",
        "each grade actually needs.",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 90


def build_grade_sheet(wb, sheet_name):
    ws = wb.create_sheet(sheet_name)
    row = 1
    for s in range(1, EXAMPLE_SUBJECTS_PER_GRADE + 1):
        ws.cell(row=row, column=1, value=f"Subject {s} (replace with real subject name)")
        row += 1
        ws.cell(row=row, column=1, value="Skill (optional grouping label)")
        row += 1
        for n in range(1, EXAMPLE_STANDARDS_PER_SUBJECT + 1):
            ws.cell(row=row, column=1, value=f"Standard {n}")
            ws.cell(row=row, column=2, value="")  # fill in the real standard text here
            row += 1
        row += 1  # blank separator row between subjects
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 90


def main():
    if len(sys.argv) < 2:
        print("usage: python3 make_blank_template.py <year-label, e.g. 2026-27> [out-path]")
        sys.exit(1)

    year_label = sys.argv[1]
    default_name = f"blank_standards_template_{year_label.replace('-', '_')}.xlsx"
    out_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "templates", default_name
    )

    wb = openpyxl.Workbook()
    build_instructions_sheet(wb)
    for sheet_name in GRADES:
        build_grade_sheet(wb, sheet_name)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    print(f"✅ Wrote blank template: {out_path}")


if __name__ == "__main__":
    main()
